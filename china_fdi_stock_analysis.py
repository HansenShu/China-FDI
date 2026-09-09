"""
China Global Investment Tracker (CGIT) — annual cumulative FDI stock proxy analysis

Uses Dataset 1 (Investments) only.

IMPORTANT INTERPRETATION
------------------------
The CGIT workbook is transaction-level investment data. It does not provide a true
FDI position/stock series because it does not systematically record divestments,
asset depreciation, write-downs, exchange-rate valuation changes, or market-price
revaluations. This script therefore constructs a GROSS CUMULATIVE INVESTMENT STOCK
PROXY by summing all recorded investment transaction values up to each year.

This is useful for studying how the accumulated geographical and sectoral footprint
of recorded Chinese overseas investment changes over time, but it should not be
labelled an official FDI stock measure.

Usage
-----
    python china_fdi_stock_analysis.py "China-Global-Investment-Tracker-public.xlsx"

Outputs
-------
- Updated workbook with annual cumulative stock-proxy tables and charts
"""

from __future__ import annotations

import argparse
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.chart import AreaChart, BarChart, LineChart, Reference
from openpyxl.drawing.line import LineProperties


SHEET_NAME = "Dataset 1"
HEADER_ROW = 5  # Excel row 6 contains column names; pandas is zero-indexed.
VALUE_COL = "Quantity in Millions"


def add_formula_analysis_sheet(workbook: Path) -> None:
    """Add annual flow and cumulative stock-proxy tables to the source workbook."""
    wb = openpyxl.load_workbook(workbook)
    source = wb[SHEET_NAME]
    analysis_name = "FDI Stock Analysis"
    if analysis_name in wb.sheetnames:
        del wb[analysis_name]
    ws = wb.create_sheet(analysis_name)

    data_start = HEADER_ROW + 2
    data_end = source.max_row
    years = sorted({
        source.cell(row, 2).value
        for row in range(data_start, data_end + 1)
        if isinstance(source.cell(row, 2).value, (int, float))
    })
    countries = sorted({
        source.cell(row, 10).value
        for row in range(data_start, data_end + 1)
        if source.cell(row, 10).value not in (None, "")
    })
    sectors = sorted({
        source.cell(row, 8).value
        for row in range(data_start, data_end + 1)
        if source.cell(row, 8).value not in (None, "")
    })

    ws["A1"] = "China FDI annual analysis from Dataset 1"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "Flows are annual sums of Quantity in Millions. Stock is a gross cumulative "
        "recorded-investment proxy, not an official FDI position measure."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(3, len(countries) + 1))

    def write_table(start_row: int, title: str, category_header: str, categories: list[str]) -> int:
        flow_row = start_row + 1
        stock_row = flow_row + len(years) + 2
        end_col = len(categories) + 1
        ws.cell(start_row, 1, title).font = Font(bold=True, size=12)
        ws.cell(flow_row, 1, "Annual flow (USD millions)").font = Font(bold=True)
        ws.cell(stock_row, 1, "Cumulative stock proxy (USD millions)").font = Font(bold=True)
        for col, category in enumerate(categories, start=2):
            ws.cell(flow_row, col, category)
            ws.cell(stock_row, col, category)
        for offset, year in enumerate(years, start=1):
            row = flow_row + offset
            ws.cell(row, 1, year)
            stock_data_row = stock_row + offset
            ws.cell(stock_data_row, 1, year)
            for col in range(2, end_col + 1):
                category_ref = f"{openpyxl.utils.get_column_letter(col)}${flow_row}"
                ws.cell(
                    row,
                    col,
                    f'=SUMIFS(\'{SHEET_NAME}\'!$E${data_start}:$E${data_end},'
                    f'\'{SHEET_NAME}\'!$B${data_start}:$B${data_end},$A{row},'
                    f'\'{SHEET_NAME}\'!${category_header}${data_start}:'
                    f'${category_header}${data_end},{category_ref})',
                )
                ws.cell(
                    stock_data_row,
                    col,
                    f'=SUMIFS(\'{SHEET_NAME}\'!$E${data_start}:$E${data_end},'
                    f'\'{SHEET_NAME}\'!$B${data_start}:$B${data_end},"<="&$A{stock_data_row},'
                    f'\'{SHEET_NAME}\'!${category_header}${data_start}:'
                    f'${category_header}${data_end},{category_ref})',
                )
        for row in range(flow_row, stock_row + len(years) + 1):
            for col in range(2, end_col + 1):
                ws.cell(row, col).number_format = '#,##0.0'
        return stock_row + len(years) + 2

    next_row = write_table(4, "By destination country", "J", countries)
    sector_start = next_row
    write_table(sector_start, "By investment sector", "H", sectors)

    def category_totals(category_col: int, categories: list[str]) -> list[str]:
        totals = {category: 0 for category in categories}
        for row in range(data_start, data_end + 1):
            category = source.cell(row, category_col).value
            value = source.cell(row, 5).value
            if category in totals and isinstance(value, (int, float)):
                totals[category] += value
        return sorted(categories, key=lambda category: totals[category], reverse=True)

    def add_stock_chart(
        anchor: str,
        title: str,
        stock_header_row: int,
        stock_first_row: int,
        categories: list[str],
        max_categories: int = 10,
    ) -> None:
        chart = LineChart()
        chart.title = title
        chart.style = 13
        chart.y_axis.title = "Cumulative recorded investment (USD millions)"
        chart.x_axis.title = "Year"
        chart.x_axis.numFmt = "0"
        chart.x_axis.tickLblPos = "low"
        chart.x_axis.majorTickMark = "out"
        chart.x_axis.delete = False
        chart.height = 9
        chart.width = 16
        years_ref = Reference(ws, min_col=1, min_row=stock_first_row, max_row=stock_first_row + len(years) - 1)
        for category in categories[:max_categories]:
            col = next(
                col for col in range(2, ws.max_column + 1)
                if ws.cell(stock_header_row, col).value == category
            )
            values = Reference(
                ws, min_col=col, min_row=stock_header_row,
                max_row=stock_first_row + len(years) - 1,
            )
            chart.add_data(values, titles_from_data=True)
        chart.set_categories(years_ref)
        chart.legend.position = "r"
        chart_colors = [
            "1F77B4", "FF7F0E", "2CA02C", "D62728", "9467BD",
            "8C564B", "E377C2", "7F7F7F", "BCBD22", "17BECF",
            "393B79", "637939", "8C6D31", "843C39",
        ]
        for index, series in enumerate(chart.ser):
            color = chart_colors[index % len(chart_colors)]
            series.graphicalProperties.line = LineProperties(solidFill=color)
            series.graphicalProperties.line.width = 24000
            series.graphicalProperties.line.noFill = False
        ws.add_chart(chart, anchor)

    country_stock_row = 4 + 1 + len(years) + 2
    sector_stock_row = sector_start + 1 + len(years) + 2
    country_ranked = category_totals(10, countries)
    sector_ranked = category_totals(8, sectors)
    add_stock_chart(
        "A100",
        "China's cumulative recorded investment by top destination countries",
        country_stock_row,
        country_stock_row + 1,
        country_ranked,
    )
    add_stock_chart(
        "J100",
        "China's cumulative recorded investment by sector",
        sector_stock_row,
        sector_stock_row + 1,
        sector_ranked,
        max_categories=len(sectors),
    )

    def write_composition_table(
        start_row: int,
        start_col: int,
        stock_header_row: int,
        stock_first_row: int,
        ranked_categories: list[str],
    ) -> tuple[int, int]:
        categories = ranked_categories[:10]
        table_categories = categories + ["Other"]
        ws.cell(start_row, start_col, "100% stacked area chart data").font = Font(bold=True, size=12)
        for offset, category in enumerate(table_categories):
            cell = ws.cell(start_row + 1, start_col + offset)
            cell.value = category
        for year_offset, year in enumerate(years):
            row = start_row + 2 + year_offset
            source_row = stock_first_row + year_offset
            ws.cell(row, start_col - 1, year)
            for category_offset, category in enumerate(categories):
                target = ws.cell(row, start_col + category_offset)
                source_col = next(
                    col for col in range(2, ws.max_column + 1)
                    if ws.cell(stock_header_row, col).value == category
                )
                target.value = f"={ws.cell(source_row, source_col).coordinate}"
            total_ref = f"SUM({openpyxl.utils.get_column_letter(start_col)}{row}:{openpyxl.utils.get_column_letter(start_col + len(categories) - 1)}{row})"
            ws.cell(
                row,
                start_col + len(categories),
                f"=SUM($B{stock_first_row + year_offset}:$"
                f"{openpyxl.utils.get_column_letter(len(ranked_categories) + 1)}{stock_first_row + year_offset})-{total_ref}",
            )
        return start_row + 1, start_row + 2

    def add_composition_chart(
        anchor: str,
        title: str,
        header_row: int,
        first_data_row: int,
        start_col: int,
        category_count: int,
    ) -> None:
        chart = AreaChart()
        chart.title = title
        chart.style = 13
        chart.grouping = "percentStacked"
        chart.overlap = 100
        chart.y_axis.title = "Share of cumulative stock (%)"
        chart.x_axis.title = "Year"
        chart.x_axis.numFmt = "0"
        chart.x_axis.tickLblPos = "low"
        chart.x_axis.delete = False
        chart.height = 9
        chart.width = 16
        data = Reference(
            ws,
            min_col=start_col,
            max_col=start_col + category_count - 1,
            min_row=header_row,
            max_row=first_data_row + len(years) - 1,
        )
        categories = Reference(ws, min_col=start_col - 1, min_row=first_data_row, max_row=first_data_row + len(years) - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.legend.position = "r"
        chart_colors = [
            "1F77B4", "FF7F0E", "2CA02C", "D62728", "9467BD",
            "8C564B", "E377C2", "7F7F7F", "BCBD22", "17BECF",
            "393B79",
        ]
        for index, series in enumerate(chart.ser):
            color = chart_colors[index % len(chart_colors)]
            series.graphicalProperties.solidFill = color
            series.graphicalProperties.line.solidFill = color
        ws.add_chart(chart, anchor)

    country_comp_header, country_comp_first = write_composition_table(
        130,
        2,
        country_stock_row,
        country_stock_row + 1,
        country_ranked,
    )
    sector_comp_header, sector_comp_first = write_composition_table(
        130, 16, sector_stock_row, sector_stock_row + 1, sector_ranked
    )
    add_composition_chart(
        "A165",
        "Changing country composition of cumulative recorded investment",
        country_comp_header,
        country_comp_first,
        2,
        11,
    )
    add_composition_chart(
        "J165",
        "Changing sector composition of cumulative recorded investment",
        sector_comp_header,
        sector_comp_first,
        16,
        11,
    )

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for row in ws.iter_rows():
        for cell in row:
            if cell.row in (5, next_row + 1) or (
                cell.value in ("Annual flow (USD millions)", "Cumulative stock proxy (USD millions)")
            ):
                cell.fill = header_fill
    ws.freeze_panes = "B6"
    ws.column_dimensions["A"].width = 34
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(workbook)


def add_canada_analysis_sheet(workbook: Path) -> None:
    """Add annual Canadian flows, cumulative stock proxy, and sector shares."""
    wb = openpyxl.load_workbook(workbook)
    source = wb[SHEET_NAME]
    analysis_name = "Canada FDI Analysis"
    if analysis_name in wb.sheetnames:
        del wb[analysis_name]
    ws = wb.create_sheet(analysis_name)

    data_start = HEADER_ROW + 2
    data_end = source.max_row
    canada_rows = [
        row for row in range(data_start, data_end + 1)
        if str(source.cell(row, 10).value).strip() == "Canada"
    ]
    years = sorted({
        source.cell(row, 2).value for row in canada_rows
        if isinstance(source.cell(row, 2).value, (int, float))
    })
    sectors = sorted({
        source.cell(row, 8).value for row in canada_rows
        if source.cell(row, 8).value not in (None, "")
    })
    if not years or not sectors:
        raise ValueError("No Canadian observations with year and sector were found.")

    ws["A1"] = "China FDI in Canada: annual sector analysis from Dataset 1"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "Annual flow is the sum of recorded transaction values. Stock is a gross "
        "cumulative recorded-investment proxy, not an official FDI position measure."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(sectors) + 2)

    flow_header = 4
    flow_first = flow_header + 1
    stock_header = flow_first + len(years) + 2
    stock_first = stock_header + 1
    total_col = len(sectors) + 2

    ws.cell(flow_header, 1, "Annual flow (USD millions)").font = Font(bold=True)
    ws.cell(stock_header, 1, "Cumulative stock proxy (USD millions)").font = Font(bold=True)
    for col, sector in enumerate(sectors, start=2):
        ws.cell(flow_header, col, sector)
        ws.cell(stock_header, col, sector)
    ws.cell(flow_header, total_col, "Total")
    ws.cell(stock_header, total_col, "Total")

    for offset, year in enumerate(years):
        flow_row = flow_first + offset
        stock_row = stock_first + offset
        ws.cell(flow_row, 1, year)
        ws.cell(stock_row, 1, year)
        for col in range(2, total_col):
            sector_ref = f"{openpyxl.utils.get_column_letter(col)}${flow_header}"
            ws.cell(
                flow_row, col,
                f'=SUMIFS(\'{SHEET_NAME}\'!$E${data_start}:$E${data_end},'
                f'\'{SHEET_NAME}\'!$B${data_start}:$B${data_end},$A{flow_row},'
                f'\'{SHEET_NAME}\'!$H${data_start}:$H${data_end},{sector_ref},'
                f'\'{SHEET_NAME}\'!$J${data_start}:$J${data_end},"Canada")',
            )
            ws.cell(
                stock_row, col,
                f"=SUM({openpyxl.utils.get_column_letter(col)}${flow_first}:"
                f"{openpyxl.utils.get_column_letter(col)}{flow_row})",
            )
        ws.cell(flow_row, total_col, f"=SUM(B{flow_row}:{openpyxl.utils.get_column_letter(total_col - 1)}{flow_row})")
        ws.cell(stock_row, total_col, f"=SUM(B{stock_row}:{openpyxl.utils.get_column_letter(total_col - 1)}{stock_row})")

    summary_header = stock_first + len(years) + 2
    ws.cell(summary_header, 1, "Latest-year sector summary").font = Font(bold=True, size=12)
    for col, heading in enumerate(
        ["Sector", "Stock proxy (USD millions)", "Share of total stock (%)", "Share change since first year (pp)"],
        start=1,
    ):
        ws.cell(summary_header + 1, col, heading)
    first_stock_row = stock_first
    latest_stock_row = stock_first + len(years) - 1
    for offset, sector in enumerate(sectors, start=summary_header + 2):
        sector_col = next(col for col in range(2, total_col) if ws.cell(stock_header, col).value == sector)
        col_letter = openpyxl.utils.get_column_letter(sector_col)
        ws.cell(offset, 1, sector)
        ws.cell(offset, 2, f"={col_letter}{latest_stock_row}")
        ws.cell(offset, 3, f"={col_letter}{latest_stock_row}/$" +
                f"{openpyxl.utils.get_column_letter(total_col)}${latest_stock_row}*100")
        ws.cell(offset, 4, f"=C{offset}-{col_letter}{first_stock_row}/$" +
                f"{openpyxl.utils.get_column_letter(total_col)}${first_stock_row}*100")

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for row in (flow_header, stock_header, summary_header + 1):
        for cell in ws[row]:
            cell.fill = header_fill
    for row in range(flow_first, stock_first + len(years)):
        for col in range(2, total_col + 1):
            ws.cell(row, col).number_format = "#,##0.0"
    for row in range(summary_header + 2, summary_header + 2 + len(sectors)):
        ws.cell(row, 2).number_format = "#,##0.0"
        ws.cell(row, 3).number_format = "0.0"
        ws.cell(row, 4).number_format = "0.0"

    chart = LineChart()
    chart.title = "China's cumulative recorded investment in Canada by sector"
    chart.y_axis.title = "USD millions"
    chart.x_axis.title = "Year"
    chart.add_data(
        Reference(ws, min_col=2, max_col=total_col - 1, min_row=stock_header, max_row=latest_stock_row),
        titles_from_data=True,
    )
    chart.set_categories(Reference(ws, min_col=1, min_row=stock_first, max_row=latest_stock_row))
    chart.height = 9
    chart.width = 16
    ws.add_chart(chart, f"F{summary_header}")

    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 38
    for col in range(2, total_col + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(workbook)


def add_energy_subsector_analysis_sheet(workbook: Path) -> None:
    """Add annual energy-subsector flows and stock proxies globally and in Canada."""
    wb = openpyxl.load_workbook(workbook)
    source = wb[SHEET_NAME]
    analysis_name = "Energy Subsector Analysis"
    if analysis_name in wb.sheetnames:
        del wb[analysis_name]
    ws = wb.create_sheet(analysis_name)

    data_start = HEADER_ROW + 2
    data_end = source.max_row
    years = sorted({
        source.cell(row, 2).value
        for row in range(data_start, data_end + 1)
        if isinstance(source.cell(row, 2).value, (int, float))
    })
    raw_subsectors = {
        source.cell(row, 9).value
        for row in range(data_start, data_end + 1)
        if source.cell(row, 8).value == "Energy"
    }
    subsectors = sorted(
        "Unspecified" if value in (None, "") else str(value).strip()
        for value in raw_subsectors
    )
    if not years or not subsectors:
        raise ValueError("No energy subsector observations were found.")

    ws["A1"] = "China energy investment: annual subsector analysis"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "Annual flow is the sum of recorded transaction values. Stock is a gross "
        "cumulative recorded-investment proxy, not an official FDI position measure."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(subsectors) + 2)

    def write_panel(start_row: int, title: str, country_criteria: str | None) -> dict[str, int]:
        flow_header = start_row + 1
        flow_first = flow_header + 1
        stock_header = flow_first + len(years) + 2
        stock_first = stock_header + 1
        total_col = len(subsectors) + 2
        ws.cell(start_row, 1, title).font = Font(bold=True, size=12)
        ws.cell(flow_header, 1, "Annual flow (USD millions)").font = Font(bold=True)
        ws.cell(stock_header, 1, "Cumulative stock proxy (USD millions)").font = Font(bold=True)
        for col, subsector in enumerate(subsectors, start=2):
            ws.cell(flow_header, col, subsector)
            ws.cell(stock_header, col, subsector)
        ws.cell(flow_header, total_col, "Total")
        ws.cell(stock_header, total_col, "Total")

        for offset, year in enumerate(years):
            flow_row = flow_first + offset
            stock_row = stock_first + offset
            ws.cell(flow_row, 1, year)
            ws.cell(stock_row, 1, year)
            for col in range(2, total_col):
                subsector_ref = f"{openpyxl.utils.get_column_letter(col)}${flow_header}"
                criteria = f'""' if ws.cell(flow_header, col).value == "Unspecified" else subsector_ref
                country_clause = ""
                if country_criteria is not None:
                    country_clause = (
                        f',\'{SHEET_NAME}\'!$J${data_start}:$J${data_end},"{country_criteria}"'
                    )
                ws.cell(
                    flow_row, col,
                    f'=SUMIFS(\'{SHEET_NAME}\'!$E${data_start}:$E${data_end},'
                    f'\'{SHEET_NAME}\'!$B${data_start}:$B${data_end},$A{flow_row},'
                    f'\'{SHEET_NAME}\'!$H${data_start}:$H${data_end},"Energy",'
                    f'\'{SHEET_NAME}\'!$I${data_start}:$I${data_end},{criteria}'
                    f'{country_clause})',
                )
                col_letter = openpyxl.utils.get_column_letter(col)
                ws.cell(
                    stock_row, col,
                    f"=SUM({col_letter}${flow_first}:{col_letter}{flow_row})",
                )
            ws.cell(
                flow_row, total_col,
                f"=SUM(B{flow_row}:{openpyxl.utils.get_column_letter(total_col - 1)}{flow_row})",
            )
            ws.cell(
                stock_row, total_col,
                f"=SUM(B{stock_row}:{openpyxl.utils.get_column_letter(total_col - 1)}{stock_row})",
            )
        return {
            "flow_header": flow_header,
            "flow_first": flow_first,
            "stock_header": stock_header,
            "stock_first": stock_first,
            "latest_stock": stock_first + len(years) - 1,
            "total_col": total_col,
        }

    global_panel = write_panel(4, "All countries", None)
    canada_start = global_panel["latest_stock"] + 4
    canada_panel = write_panel(canada_start, "Canada only", "Canada")

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for panel in (global_panel, canada_panel):
        for row in (panel["flow_header"], panel["stock_header"]):
            for cell in ws[row]:
                cell.fill = header_fill
        for row in range(panel["flow_first"], panel["latest_stock"] + 1):
            for col in range(2, panel["total_col"] + 1):
                ws.cell(row, col).number_format = "#,##0.0"

    def add_stock_chart(panel: dict[str, int], anchor: str, title: str) -> None:
        chart = LineChart()
        chart.title = title
        chart.y_axis.title = "Cumulative recorded investment (USD millions)"
        chart.x_axis.title = "Year"
        chart.x_axis.numFmt = "0"
        chart.x_axis.tickLblPos = "low"
        chart.x_axis.majorTickMark = "out"
        chart.x_axis.delete = False
        chart.add_data(
            Reference(
                ws,
                min_col=2,
                max_col=panel["total_col"] - 1,
                min_row=panel["stock_header"],
                max_row=panel["latest_stock"],
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(ws, min_col=1, min_row=panel["stock_first"], max_row=panel["latest_stock"])
        )
        chart.height = 9
        chart.width = 16
        ws.add_chart(chart, anchor)

    add_stock_chart(
        global_panel, f"J{global_panel['flow_header']}",
        "Cumulative energy investment by subsector: all countries",
    )
    add_stock_chart(
        canada_panel, f"J{canada_panel['flow_header']}",
        "Cumulative energy investment by subsector: Canada",
    )

    ws.freeze_panes = "B6"
    ws.column_dimensions["A"].width = 38
    for col in range(2, len(subsectors) + 3):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 18
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(workbook)


def add_deal_concentration_analysis_sheet(
    workbook: Path, df: pd.DataFrame
) -> None:
    """Add annual deal-size and concentration metrics for Canada and globally."""
    wb = openpyxl.load_workbook(workbook)
    helper_name = "Deal Concentration Calc"
    if helper_name in wb.sheetnames:
        del wb[helper_name]
    helper = wb.create_sheet(helper_name)
    analysis_name = "Deal Size & Concentration"
    if analysis_name in wb.sheetnames:
        del wb[analysis_name]
    ws = wb.create_sheet(analysis_name)

    data_start = HEADER_ROW + 2
    source_end = HEADER_ROW + 1 + len(df)
    years = list(range(int(df["Year"].min()), int(df["Year"].max()) + 1))
    helper["A1"] = "Year"
    helper["B1"] = "Country"
    helper["C1"] = "Value (USD mn)"
    for row in range(2, len(df) + 2):
        source_row = data_start + row - 2
        helper.cell(row, 1, f"='{SHEET_NAME}'!B{source_row}")
        helper.cell(row, 2, f"='{SHEET_NAME}'!J{source_row}")
        helper.cell(row, 3, f"='{SHEET_NAME}'!E{source_row}")

    helper_columns: dict[tuple[str, int], str] = {}
    for index, year in enumerate(years, start=4):
        global_col = openpyxl.utils.get_column_letter(index)
        canada_col = openpyxl.utils.get_column_letter(index + len(years))
        helper.cell(1, index, f"All countries {year}")
        helper.cell(1, index + len(years), f"Canada {year}")
        for row in range(2, len(df) + 2):
            helper.cell(row, index, f'=IF($A{row}={year},$C{row},"")')
            helper.cell(row, index + len(years), f'=IF(AND($A{row}={year},$B{row}="Canada"),$C{row},"")')
        helper_columns[("All countries", year)] = global_col
        helper_columns[("Canada", year)] = canada_col
    helper.sheet_state = "hidden"

    metric_names = [
        "Total recorded investment (USD mn)",
        "Transactions",
        "Mean deal size (USD mn)",
        "Median deal size (USD mn)",
        "Largest deal (USD mn)",
        "90th percentile share (%)",
        "Excluding largest deal (USD mn)",
    ]

    ws["A1"] = "Deal size and concentration of China's recorded investment"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        "Canada is compared with all countries. Concentration measures describe recorded "
        "transactions, not official FDI positions. The 90th-percentile share is the "
        "share of annual investment from transactions at or above that year's 90th-percentile deal size."
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=12)

    header_row = 4
    ws.cell(header_row, 1, "Year").font = Font(bold=True)
    columns = [
        ("Canada", name) for name in metric_names
    ] + [("All countries", name) for name in metric_names]
    for col, (group, name) in enumerate(columns, start=2):
        ws.cell(header_row, col, f"{group}: {name}")

    for row_offset, year in enumerate(years, start=1):
        row = header_row + row_offset
        ws.cell(row, 1, year)
        for col, (group, name) in enumerate(columns, start=2):
            helper_col = helper_columns[(group, year)]
            helper_range = f"'{helper_name}'!${helper_col}$2:${helper_col}${len(df) + 1}"
            total = f"SUM({helper_range})"
            formulas = {
                "Total recorded investment (USD mn)": f'=IFERROR({total},0)',
                "Transactions": f'=COUNT({helper_range})',
                "Mean deal size (USD mn)": f'=IFERROR(AVERAGE({helper_range}),0)',
                "Median deal size (USD mn)": f'=IFERROR(MEDIAN({helper_range}),0)',
                "Largest deal (USD mn)": f'=IFERROR(MAX({helper_range}),0)',
                "90th percentile share (%)": (
                    f'=IFERROR(SUMIF({helper_range},">="&PERCENTILE({helper_range},0.90),'
                    f'{helper_range})/{total}*100,0)'
                ),
                "Excluding largest deal (USD mn)": f'=IFERROR({total}-MAX({helper_range}),0)',
            }
            ws.cell(row, col, formulas[name])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[header_row]:
        cell.fill = header_fill
    for row in range(header_row + 1, header_row + len(years) + 1):
        for col in range(2, 16):
            ws.cell(row, col).number_format = "#,##0.0"
    for col in range(2, 16):
        if "share" in str(ws.cell(header_row, col).value).lower():
            for row in range(header_row + 1, header_row + len(years) + 1):
                ws.cell(row, col).number_format = "0.0"

    def add_chart(
        anchor: str,
        title: str,
        selected_columns: list[int],
        y_title: str,
        min_row: int = header_row,
    ) -> None:
        chart = LineChart()
        chart.title = title
        chart.y_axis.title = y_title
        chart.x_axis.title = "Year"
        chart.x_axis.numFmt = "0"
        chart.x_axis.tickLblPos = "low"
        chart.x_axis.majorTickMark = "out"
        chart.x_axis.delete = False
        for col in selected_columns:
            chart.add_data(
                Reference(
                    ws,
                    min_col=col,
                    max_col=col,
                    min_row=min_row,
                    max_row=header_row + len(years),
                ),
                titles_from_data=True,
            )
        chart.set_categories(
            Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + len(years))
        )
        chart.height = 8
        chart.width = 15
        chart.legend.position = "r"
        ws.add_chart(chart, anchor)

    def add_activity_chart(anchor: str) -> None:
        bars = BarChart()
        bars.type = "col"
        bars.title = "Canada: total investment and transaction count"
        bars.y_axis.title = "Total investment (USD millions)"
        bars.x_axis.title = "Year"
        bars.x_axis.numFmt = "0"
        bars.x_axis.tickLblPos = "low"
        bars.x_axis.majorTickMark = "out"
        bars.x_axis.delete = False
        bars.add_data(
            Reference(ws, min_col=2, max_col=2, min_row=header_row, max_row=header_row + len(years)),
            titles_from_data=True,
        )
        bars.set_categories(
            Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + len(years))
        )
        bars.height = 8
        bars.width = 15

        transactions = LineChart()
        transactions.y_axis.axId = 200
        transactions.y_axis.title = "Transactions"
        transactions.y_axis.crosses = "max"
        transactions.add_data(
            Reference(ws, min_col=3, max_col=3, min_row=header_row, max_row=header_row + len(years)),
            titles_from_data=True,
        )
        transactions.set_categories(
            Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + len(years))
        )
        bars += transactions
        ws.add_chart(bars, anchor)

    # Canada columns: total=2, transactions=3, mean=4, median=5, p90 share=7.
    add_activity_chart("W4")
    add_chart(
        "W20",
        "Canada: mean versus median deal size",
        [4, 5],
        "USD millions",
    )
    add_chart(
        "W52",
        "90th-percentile deal concentration: Canada versus all countries",
        [7, 14],
        "Share of annual investment (%)",
    )

    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 12
    for col in range(2, 16):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 19
    for col in range(1, 3 + len(years) * 2):
        helper.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 15
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(workbook)


def add_deal_concentration_analysis_sheet(
    workbook: Path, df: pd.DataFrame
) -> None:
    """Rebuild the average-versus-median transaction comparison tab."""
    wb = openpyxl.load_workbook(workbook)
    helper_name = "Deal Concentration Calc"
    if helper_name in wb.sheetnames:
        del wb[helper_name]
    sheet_name = "Deal Size & Concentration"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    data_start = HEADER_ROW + 2
    source_end = HEADER_ROW + 1 + len(df)
    years = list(range(int(df["Year"].min()), int(df["Year"].max()) + 1))

    ws["A1"] = "Average and median transaction size over time"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Transaction values are in USD millions. Rest of world excludes Canada."
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=15)

    metrics = ["Mean transaction size (USD mn)", "Median transaction size (USD mn)"]
    ws.cell(4, 1, "Year")
    groups = ["Canada", "Rest of world"]
    for col, (group, metric) in enumerate(
        [(g, m) for g in groups for m in metrics], start=2
    ):
        ws.cell(4, col, f"{group}: {metric}")
    for row_offset, year in enumerate(years, start=1):
        row = 4 + row_offset
        ws.cell(row, 1, year)
        for col, (group, metric) in enumerate(
            [(g, m) for g in groups for m in metrics], start=2
        ):
            year_range = f"'{SHEET_NAME}'!$B${data_start}:$B${source_end}"
            value_range = f"'{SHEET_NAME}'!$E${data_start}:$E${source_end}"
            country_range = f"'{SHEET_NAME}'!$J${data_start}:$J${source_end}"
            country_condition = (
                f',{country_range},"Canada"' if group == "Canada"
                else f',{country_range},"<>"&"Canada"'
            )
            criteria = f"{year_range},$A{row}{country_condition}"
            formulas = {
                "Mean transaction size (USD mn)": (
                    f'=IF(COUNTIFS({year_range},$A{row},{country_range},'
                    f'{"\"Canada\"" if group == "Canada" else "\"<>\"&\"Canada\""})=0,"",'
                    f'AVERAGEIFS({value_range},{criteria}))'
                ),
                "Median transaction size (USD mn)": (
                    f'=IFERROR(MEDIAN(_xlfn.FILTER({value_range},'
                    f'({year_range}=$A{row})*'
                    f'({country_range}{"=" if group == "Canada" else "<>"}"Canada"))),"")'
                ),
            }
            ws.cell(row, col, formulas[metric])

    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[4]:
        cell.fill = fill
    for row in range(5, 5 + len(years)):
        for col in range(2, 6):
            ws.cell(row, col).number_format = "0.0"
    for col in (2, 3, 4, 5):
        for row in range(5, 5 + len(years)):
            ws.cell(row, col).number_format = "0.0"

    def line_chart(anchor, title, columns, y_title):
        chart = LineChart()
        chart.title = title
        chart.y_axis.title = y_title
        chart.x_axis.title = "Year"
        chart.x_axis.numFmt = "0"
        chart.x_axis.delete = False
        for col in columns:
            chart.add_data(
                Reference(ws, min_col=col, max_col=col, min_row=4, max_row=4 + len(years)),
                titles_from_data=True,
            )
        chart.set_categories(
            Reference(ws, min_col=1, min_row=5, max_row=4 + len(years))
        )
        chart.height = 8
        chart.width = 15
        ws.add_chart(chart, anchor)

    line_chart("Q4", "Canada: mean versus median transaction size", [2, 3], "USD millions")
    line_chart(
        "Q20",
        "Rest of world: mean versus median transaction size",
        [4, 5],
        "USD millions",
    )
    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 12
    for col in range(2, 6):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 19
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.save(workbook)


def load_and_clean(workbook: Path) -> pd.DataFrame:
    """Load Dataset 1 and retain valid annual investment observations."""
    df = pd.read_excel(workbook, sheet_name=SHEET_NAME, header=HEADER_ROW)

    required = {"Year", VALUE_COL, "Country", "Sector"}
    missing = required.difference(df.columns)
    if missing:
        raise ValueError(f"Missing required columns in {SHEET_NAME}: {sorted(missing)}")

    out = df.copy()
    out["Year"] = pd.to_numeric(out["Year"], errors="coerce")
    out[VALUE_COL] = pd.to_numeric(out[VALUE_COL], errors="coerce")

    # A transaction needs a year and a reported investment value to enter the stock proxy.
    out = out.dropna(subset=["Year", VALUE_COL]).copy()
    out["Year"] = out["Year"].astype(int)

    # Keep unknown geography/sector rather than silently dropping transaction values.
    out["Country"] = out["Country"].fillna("Unspecified").astype(str).str.strip()
    out["Sector"] = out["Sector"].fillna("Unspecified").astype(str).str.strip()

    # Negative values, if ever present, would imply reductions; CGIT is mainly gross deals.
    # We preserve them rather than force them to zero.
    return out


def annual_panel(df: pd.DataFrame, category: str) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Return annual flows, cumulative stock proxy, and cumulative-stock shares."""
    years = np.arange(df["Year"].min(), df["Year"].max() + 1)

    flows = (
        df.groupby(["Year", category], as_index=False)[VALUE_COL]
        .sum()
        .pivot(index="Year", columns=category, values=VALUE_COL)
        .reindex(years, fill_value=0)
        .fillna(0)
        .sort_index()
    )

    stock = flows.cumsum()
    total_stock = stock.sum(axis=1)
    shares = stock.div(total_stock.replace(0, np.nan), axis=0) * 100

    return flows, stock, shares


def latest_summary(stock: pd.DataFrame, shares: pd.DataFrame) -> pd.DataFrame:
    """Rank categories by cumulative stock proxy in the latest year."""
    latest_year = int(stock.index.max())
    summary = pd.DataFrame(
        {
            "Stock proxy (USD millions)": stock.loc[latest_year],
            "Share of total stock (%)": shares.loc[latest_year],
        }
    )
    summary.index.name = "Category"
    summary = summary.sort_values("Stock proxy (USD millions)", ascending=False)
    summary.insert(0, "Year", latest_year)
    return summary


def top_categories(stock: pd.DataFrame, top_n: int) -> list[str]:
    """Choose top categories based on latest-year cumulative stock."""
    return stock.loc[stock.index.max()].nlargest(top_n).index.tolist()


def save_line_chart(stock: pd.DataFrame, categories: list[str], title: str,
                    ylabel: str, output: Path) -> None:
    fig, ax = plt.subplots(figsize=(11, 6.5))
    for category in categories:
        ax.plot(stock.index, stock[category] / 1000, linewidth=2, label=category)
    ax.set_title(title)
    ax.set_xlabel("Year")
    ax.set_ylabel(ylabel)
    ax.grid(alpha=0.25)
    ax.legend(frameon=False, ncol=2)
    fig.tight_layout()
    fig.savefig(output, dpi=180, bbox_inches="tight")
    plt.close(fig)


def save_stacked_share_chart(shares: pd.DataFrame, categories: list[str], title: str,
                             output: Path) -> None:
    # Top categories are shown individually; everything else is combined into "Other".
    plot_df = shares[categories].copy()
    plot_df["Other"] = 100 - plot_df.sum(axis=1)

    fig, ax = plt.subplots(figsize=(11, 6.5))
    ax.stackplot(plot_df.index, [plot_df[c] for c in plot_df.columns],
                 labels=plot_df.columns, alpha=0.9)
    ax.set_title(title)
    ax.set_xlabel("Year")
    ax.set_ylabel("Share of cumulative stock proxy (%)")
    ax.set_ylim(0, 100)
    ax.grid(axis="y", alpha=0.25)
    ax.legend(frameon=False, ncol=2, loc="upper left", bbox_to_anchor=(1.01, 1.0))
    fig.tight_layout()
    fig.savefig(output, dpi=180, bbox_inches="tight")
    plt.close(fig)


def save_latest_bar(summary: pd.DataFrame, top_n: int, title: str, output: Path) -> None:
    top = summary.head(top_n).sort_values("Stock proxy (USD millions)")
    fig, ax = plt.subplots(figsize=(10, 6.5))
    ax.barh(top.index, top["Stock proxy (USD millions)"] / 1000)
    ax.set_title(title)
    ax.set_xlabel("Cumulative recorded investment (USD billions)")
    ax.set_ylabel("")
    ax.grid(axis="x", alpha=0.25)
    fig.tight_layout()
    fig.savefig(output, dpi=180, bbox_inches="tight")
    plt.close(fig)


def save_total_stock_chart(df: pd.DataFrame, output: Path) -> pd.DataFrame:
    annual_flow = df.groupby("Year")[VALUE_COL].sum().sort_index()
    years = np.arange(df["Year"].min(), df["Year"].max() + 1)
    annual_flow = annual_flow.reindex(years, fill_value=0)
    cumulative = annual_flow.cumsum()

    total = pd.DataFrame({
        "Annual recorded investment flow (USD millions)": annual_flow,
        "Cumulative stock proxy (USD millions)": cumulative,
    })
    total.index.name = "Year"

    fig, ax = plt.subplots(figsize=(10.5, 6))
    ax.plot(total.index, total["Cumulative stock proxy (USD millions)"] / 1000, linewidth=2.5)
    ax.set_title("China's cumulative recorded overseas investment stock proxy")
    ax.set_xlabel("Year")
    ax.set_ylabel("USD billions")
    ax.grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(output, dpi=180, bbox_inches="tight")
    plt.close(fig)

    return total


def save_change_table(shares: pd.DataFrame, stock: pd.DataFrame) -> pd.DataFrame:
    """Show how category importance changed from the first to latest year."""
    first_year = int(stock.index.min())
    latest_year = int(stock.index.max())

    result = pd.DataFrame({
        f"Stock {first_year} (USD mn)": stock.loc[first_year],
        f"Stock {latest_year} (USD mn)": stock.loc[latest_year],
        f"Share {first_year} (%)": shares.loc[first_year],
        f"Share {latest_year} (%)": shares.loc[latest_year],
    })
    result["Change in share (percentage points)"] = (
        result[f"Share {latest_year} (%)"] - result[f"Share {first_year} (%)"]
    )
    return result.sort_values(f"Stock {latest_year} (USD mn)", ascending=False)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Analyze annual cumulative Chinese overseas investment using CGIT Dataset 1."
    )
    default_workbook = Path(__file__).resolve().parent / "China-Global-Investment-Tracker-public.xlsx"
    parser.add_argument(
        "workbook",
        type=Path,
        nargs="?",
        default=default_workbook,
        help=f"Path to the CGIT .xlsx workbook (default: {default_workbook.name})",
    )
    args = parser.parse_args()

    if not args.workbook.exists():
        raise FileNotFoundError(args.workbook)
    add_formula_analysis_sheet(args.workbook)
    add_canada_analysis_sheet(args.workbook)
    add_energy_subsector_analysis_sheet(args.workbook)

    df = load_and_clean(args.workbook)
    canada_df = df[df["Country"].eq("Canada")].copy()
    if canada_df.empty:
        raise ValueError("No Canadian observations were found in Dataset 1.")
    add_deal_concentration_analysis_sheet(args.workbook, df)

    _, country_stock, country_shares = annual_panel(df, "Country")
    _, sector_stock, sector_shares = annual_panel(df, "Sector")
    _, canada_sector_stock, canada_sector_shares = annual_panel(
        canada_df, "Sector"
    )

    country_summary = latest_summary(country_stock, country_shares)
    sector_summary = latest_summary(sector_stock, sector_shares)

    latest_year = int(df["Year"].max())
    first_year = int(df["Year"].min())
    latest_total_bn = country_stock.loc[latest_year].sum() / 1000
    canada_latest_year = int(canada_sector_stock.index.max())
    canada_latest_total_bn = (
        canada_sector_stock.loc[canada_latest_year].sum() / 1000
    )

    print("\nCGIT annual cumulative investment analysis complete")
    print("---------------------------------------------------")
    print(f"Years: {first_year}–{latest_year}")
    print(f"Transactions used: {len(df):,}")
    print(f"Latest cumulative stock proxy: USD {latest_total_bn:,.1f} billion")
    print(f"Top country in {latest_year}: {country_summary.index[0]}")
    print(f"Top sector in {latest_year}: {sector_summary.index[0]}")
    print(
        f"Canada: {int(canada_df['Year'].min())}–{canada_latest_year}, "
        f"{len(canada_df):,} transactions, latest stock proxy "
        f"USD {canada_latest_total_bn:,.1f} billion"
    )
    print(f"Top Canadian sector in {canada_latest_year}: {canada_sector_stock.loc[canada_latest_year].idxmax()}")
    print(f"Updated workbook: {args.workbook.resolve()}")
    print("\nNOTE: This is a gross cumulative investment stock proxy, not an official FDI position stock.")


if __name__ == "__main__":
    main()
